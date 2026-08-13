"""
Экспорт пользователей в CSV и Excel
"""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def users_to_csv(users: list, server_name: str) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([
        "password", "comment", "active", "online", "expires",
        "total_gb", "traffic_used", "max_devices", "devices_bound",
        "max_down_mbps", "max_up_mbps", "server",
    ])
    for u in users:
        writer.writerow([
            u.get("password", ""),
            u.get("comment", ""),
            u.get("active", True),
            u.get("online", False),
            u.get("expires", ""),
            u.get("total_gb", 0),
            u.get("traffic_used_fmt", ""),
            u.get("max_devices", 1),
            u.get("devices_bound", 0),
            u.get("max_down_mbps", 0),
            u.get("max_up_mbps", 0),
            server_name,
        ])
    return buf.getvalue().encode("utf-8-sig")


def users_to_xlsx(users: list, server_name: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users"

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(color="FFFFFF", bold=True)

    headers = [
        "Пароль", "Комментарий", "Активен", "Онлайн", "Истекает",
        "Трафик лимит", "Трафик использовано", "Устройств макс", "Устройств привязано",
        "Max Down Mbps", "Max Up Mbps", "Сервер",
    ]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    green_fill = PatternFill("solid", fgColor="C6EFCE")
    red_fill = PatternFill("solid", fgColor="FFC7CE")

    for row_idx, u in enumerate(users, 2):
        active = u.get("active", True)
        online = u.get("online", False)
        expires = u.get("expires", "")

        row_data = [
            u.get("password", ""),
            u.get("comment", ""),
            "Да" if active else "Нет",
            "Да" if online else "Нет",
            expires,
            u.get("total_gb", 0),
            u.get("traffic_used_fmt", ""),
            u.get("max_devices", 1),
            u.get("devices_bound", 0),
            u.get("max_down_mbps", 0),
            u.get("max_up_mbps", 0),
            server_name,
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if online:
                cell.fill = green_fill
            if not active:
                if col_idx == 3:
                    cell.fill = red_fill

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Info")
    ws2["A1"] = "Сервер"
    ws2["B1"] = server_name
    ws2["A2"] = "Экспортировано"
    ws2["B2"] = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    ws2["A3"] = "Пользователей"
    ws2["B3"] = len(users)
    ws2["A4"] = "Онлайн"
    ws2["B4"] = sum(1 for u in users if u.get("online", False))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
